#!/usr/bin/env python3
"""
AcoustEEEcare v8.4 — Host-side BLE receiver (Data Gathering Edition)
=====================================================================
Matches firmware v8.4 (BLE-only, ONE capture → BOTH HR and RR computed
on-device, dynamic MTU, no SD card, no MFCC stream).

Protocol (v8.4)
───────────────
    Device → host, once after MTU exchange:
        "READY\n"                         (gate the Record action)

    Host → device:
        "REC"                              (no organ select — always dual)

    Device → host, per recording:
        "START:<audio_bytes>:DUAL\n"
        [seq u16 LE][len u16 LE][int16 PCM]   — audio chunks (for the WAV)
        "finished\n"
        "INF:START\n"                       (on-device inference begins)
        "HR:<bpm>\n"                        (heart result, or ERR:HEART_INF)
        "RR:<bpm>\n"                        (lung result,  or ERR:LUNG_INF)
        "INF:DONE\n"                        (session complete)

    Any time while connected:
        "PING\n"                            (keepalive — ignored here)

    Fatal capture error:
        "ERR:SAADC\n" / "ERR:DSP\n" / etc.

Folder / file layout
────────────────────
    <SUBJECT_CODE>/                        ← per-subject folder (uppercase)
        <SUBJECT_CODE>_log.xlsx            ← HR/RR log for this subject
        <SUBJECT_CODE>_<ACTIVITY>_<N:03d>_<TIMESTAMP>.wav
        <SUBJECT_CODE>_<ACTIVITY>_<N:03d>_<TIMESTAMP>.txt

Usage
─────
    pip install bleak openpyxl
    python receiver_v84.py
"""

import asyncio
import struct
import wave
import time
import sys
from datetime import datetime
from pathlib import Path

from bleak import BleakScanner, BleakClient

# ── BLE / device constants ─────────────────────────────────────────────────────

NUS_TX_CHAR_UUID = "6e400003-b5a3-f393-e0a9-e50e24dcca9e"
NUS_RX_CHAR_UUID = "6e400002-b5a3-f393-e0a9-e50e24dcca9e"

DEVICE_NAME   = "AcoustEEEcare"
SAMPLE_RATE   = 8000
SCAN_TIMEOUT  = 20.0
REC_TIMEOUT   = 90.0   # 10 s record + audio drain + ~20-35 s dual inference + margin
READY_TIMEOUT = 5.0    # wait for firmware READY after connecting

VALID_ACTIVITIES = {"REST", "WALKING"}

idle_keepalive = True   # True when at the prompt (link idle)


# ── Session metadata (set at startup) ─────────────────────────────────────────

subject_code: str = ""   # uppercase, e.g. "S001"
activity:     str = ""   # "REST" or "WALKING"
output_dir:   Path = Path(".")


def prompt_session_info() -> None:
    """Ask the user for subject code and activity before connecting."""
    global subject_code, activity, output_dir

    print("=" * 55)
    print("  AcoustEEEcare — Data Gathering Setup")
    print("=" * 55)

    while True:
        code = input("  Subject code: ").strip().upper()
        if code:
            subject_code = code
            break
        print("  ✗ Subject code cannot be empty.")

    while True:
        act = input("  Activity [REST / WALKING]: ").strip().upper()
        if act in VALID_ACTIVITIES:
            activity = act
            break
        print(f"  ✗ Please enter one of: {', '.join(sorted(VALID_ACTIVITIES))}")

    output_dir = Path(subject_code)
    output_dir.mkdir(exist_ok=True)

    print()
    print(f"  Subject : {subject_code}")
    print(f"  Activity: {activity}")
    print(f"  Folder  : {output_dir}/")
    print("=" * 55)
    print()


# ── Excel log ─────────────────────────────────────────────────────────────────

def _excel_log_path() -> Path:
    return output_dir / f"{subject_code}_log.xlsx"


def _init_excel_log() -> None:
    """Create the Excel log for this subject if it does not yet exist."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    path = _excel_log_path()
    if path.exists():
        return   # already initialised; rows will be appended later

    wb = Workbook()
    ws = wb.active
    ws.title = subject_code

    # ── Header row ──
    headers = [
        "Recording #", "Activity", "Timestamp",
        "HR (BPM)", "RR (breaths/min)",
        "Audio drops", "Loss %", "Notes",
    ]
    header_font  = Font(name="Arial", bold=True, color="FFFFFF")
    header_fill  = PatternFill("solid", start_color="2E4057")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_side    = Side(style="thin", color="CCCCCC")
    thin_border  = Border(left=thin_side, right=thin_side,
                          bottom=thin_side, top=thin_side)

    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = center_align
        cell.border    = thin_border

    # Column widths
    widths = [14, 12, 22, 12, 20, 14, 10, 30]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[
            ws.cell(row=1, column=col).column_letter
        ].width = w

    ws.freeze_panes = "A2"
    wb.save(path)


def _append_excel_row(rec_num: int, hr, rr,
                      audio_drops: int, loss_pct: float,
                      notes: str = "") -> None:
    """Append one result row to the subject's Excel log."""
    from openpyxl import load_workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    path = _excel_log_path()
    wb   = load_workbook(path)
    ws   = wb.active

    row = ws.max_row + 1

    def _val(v):
        if v is None: return "n/a"
        if v == "ERR": return "ERROR"
        return round(float(v), 1)

    thin_side   = Side(style="thin", color="CCCCCC")
    thin_border = Border(left=thin_side, right=thin_side,
                         bottom=thin_side, top=thin_side)
    center      = Alignment(horizontal="center")

    row_fill = PatternFill("solid", start_color="F0F4F8") \
               if row % 2 == 0 else PatternFill("solid", start_color="FFFFFF")
    std_font  = Font(name="Arial", size=10)

    values = [
        rec_num,
        activity,
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        _val(hr),
        _val(rr),
        audio_drops,
        round(loss_pct, 2),
        notes,
    ]

    for col, val in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=val)
        cell.font      = std_font
        cell.border    = thin_border
        cell.fill      = row_fill
        cell.alignment = center if col != 8 else Alignment(horizontal="left")

    wb.save(path)


# ── Audio packet / loss tracker ────────────────────────────────────────────────

class PacketStats:
    def __init__(self):
        self.reset()

    def reset(self):
        self.audio_pkts_rx    = 0
        self.audio_bytes_rx   = 0
        self.audio_bytes_exp  = 0
        self.audio_pkts_lost  = 0
        self.audio_bytes_lost = 0
        self.audio_overhead   = 0
        self.total_notifications = 0
        self.total_raw_bytes     = 0
        self.session_start_time  = None
        self.session_end_time    = None

    def start_session(self):
        self.session_start_time = time.monotonic()

    def end_session(self):
        self.session_end_time = time.monotonic()

    @property
    def session_duration(self):
        if self.session_start_time and self.session_end_time:
            return self.session_end_time - self.session_start_time
        return 0.0

    @property
    def audio_loss_pct(self):
        total = self.audio_pkts_rx + self.audio_pkts_lost
        return self.audio_pkts_lost * 100.0 / max(total, 1)

    @property
    def audio_throughput_kbps(self):
        return (self.audio_bytes_rx * 8) / max(self.session_duration, 0.001) / 1000.0


pstats = PacketStats()


# ── Receiver state ─────────────────────────────────────────────────────────────

class StreamMode:
    IDLE  = "idle"
    AUDIO = "audio"


class State:
    def __init__(self):
        self.reset()

    def reset(self):
        self.error          = None
        self.mode           = StreamMode.IDLE
        self.expected_bytes = 0
        self.audio_done     = False
        self.audio_samples  = bytearray()
        self.audio_seq      = 0
        self.audio_gaps     = 0
        self.audio_chunks   = 0
        self.inf_started    = False
        self.inf_done       = False
        self.hr             = None
        self.rr             = None
        self.last_progress  = 0.0
        pstats.reset()
        pstats.start_session()


state       = State()
ready_event = asyncio.Event()


# ── Packet classifier ──────────────────────────────────────────────────────────

def _is_data_chunk(data: bytearray) -> bool:
    if state.mode != StreamMode.AUDIO or len(data) < 5:
        return False
    payload_len = struct.unpack_from("<H", data, 2)[0]
    return payload_len > 0 and payload_len == len(data) - 4


def handle_notification(_sender, data: bytearray):
    pstats.total_notifications += 1
    pstats.total_raw_bytes     += len(data)
    if _is_data_chunk(data):
        _handle_audio_chunk(data)
    else:
        _handle_text(data)


# ── Control / text message handler ────────────────────────────────────────────

def _handle_text(data: bytearray):
    try:
        text = data.decode("ascii", errors="replace").strip()
    except Exception:
        return

    if text == "PING":
        return

    if text.startswith("READY"):
        ready_event.set()
        print("  [READY ] Firmware ready.")
        return

    if text.startswith("START:"):
        try:
            parts = text.split(":")
            n = int(parts[1])
            state.expected_bytes   = n
            state.audio_samples    = bytearray()
            state.audio_seq        = 0
            state.audio_gaps       = 0
            state.audio_chunks     = 0
            state.audio_done       = False
            state.mode             = StreamMode.AUDIO
            pstats.audio_bytes_exp = n
            dur = n / 2 / SAMPLE_RATE
            print(f"\n  [AUDIO ] Receiving — {n} B ({n // 2} samples, {dur:.1f} s expected)")
        except (ValueError, IndexError):
            print(f"  [WARN  ] Malformed START: {text!r}")
        return

    if text.startswith("finished"):
        state.audio_done = True
        got = len(state.audio_samples)
        exp = state.expected_bytes
        pct = min(got * 100 // max(exp, 1), 100)
        drops = pstats.audio_pkts_lost
        loss  = pstats.audio_loss_pct
        drop_str = "✓ No drops" if drops == 0 else f"✗ {drops} pkt(s) lost ({loss:.2f}% loss)"
        print(f"\n  [AUDIO ] Done — {got}/{exp} B ({pct}%)  |  {drop_str}")
        state.mode = StreamMode.IDLE
        return

    if text.startswith("INF:START"):
        state.inf_started = True
        print("  [INFER ] Running on-device inference (~20–35 s)…")
        return

    if text.startswith("INF:DONE"):
        state.inf_done = True
        return

    if text.startswith("HR:"):
        try:
            state.hr = float(text.split(":")[1])
            print(f"  [HR    ] {state.hr:.1f} BPM")
        except (ValueError, IndexError):
            print(f"  [WARN  ] Malformed HR: {text!r}")
        return

    if text.startswith("RR:"):
        try:
            state.rr = float(text.split(":")[1])
            print(f"  [RR    ] {state.rr:.1f} breaths/min")
        except (ValueError, IndexError):
            print(f"  [WARN  ] Malformed RR: {text!r}")
        return

    if text.startswith("ERR:HEART_INF"):
        state.hr = "ERR"
        print("  [ERROR ] Heart inference failed on device.")
        return

    if text.startswith("ERR:LUNG_INF"):
        state.rr = "ERR"
        print("  [ERROR ] Lung inference failed on device.")
        return

    if text.startswith("ERR:"):
        state.error      = text
        state.audio_done = True
        state.inf_done   = True
        state.mode       = StreamMode.IDLE
        print(f"\n  [ERROR ] Firmware reported: {text}")
        return

    # Unrecognised firmware message — silently ignore to keep output clean


# ── Audio chunk handler ────────────────────────────────────────────────────────

def _handle_audio_chunk(data: bytearray):
    seq     = struct.unpack_from("<H", data, 0)[0]
    ln      = struct.unpack_from("<H", data, 2)[0]
    payload = data[4: 4 + ln]

    pstats.audio_pkts_rx  += 1
    pstats.audio_bytes_rx += ln
    pstats.audio_overhead += 4

    if seq != state.audio_seq:
        gap_pkts  = (seq - state.audio_seq) & 0xFFFF
        gap_bytes = gap_pkts * ln
        state.audio_samples     += b"\x00\x00" * (gap_bytes // 2)
        state.audio_gaps        += gap_pkts
        pstats.audio_pkts_lost  += gap_pkts
        pstats.audio_bytes_lost += gap_bytes

    state.audio_samples += payload
    state.audio_seq      = (seq + 1) & 0xFFFF
    state.audio_chunks  += 1
    _print_audio_progress()


def _print_audio_progress():
    now = time.monotonic()
    if now - state.last_progress < 0.25:
        return
    state.last_progress = now
    got = len(state.audio_samples)
    exp = state.expected_bytes
    pct = min(got * 100 // max(exp, 1), 100)
    bar = "█" * (pct // 5) + "░" * (20 - pct // 5)
    print(f"\r  [{bar}] {pct:3d}%  {got:>7}/{exp} B  "
          f"lost={pstats.audio_pkts_lost}  ({pstats.audio_loss_pct:.1f}% loss)",
          end="", flush=True)


# ── Save helpers ───────────────────────────────────────────────────────────────

def _base_name(rec_num: int, ts: str) -> str:
    return f"{subject_code}_{activity}_{rec_num:03d}_{ts}"


def save_wav(samples: bytearray, filename: Path) -> None:
    target = state.expected_bytes
    if len(samples) > target:
        samples = samples[:target]
    elif len(samples) < target:
        samples += b"\x00" * (target - len(samples))

    with wave.open(str(filename), "wb") as wf:
        wf.setnchannels(1)
        wf.setsampwidth(2)
        wf.setframerate(SAMPLE_RATE)
        wf.writeframes(samples)


def _fmt_result(v):
    if v is None:  return "n/a"
    if v == "ERR": return "ERROR"
    return f"{v:.1f}"


def save_results_txt(filename: Path) -> None:
    with open(filename, "w") as f:
        f.write(f"subject   : {subject_code}\n")
        f.write(f"activity  : {activity}\n")
        f.write(f"timestamp : {datetime.now().isoformat()}\n")
        f.write(f"HR_bpm    : {_fmt_result(state.hr)}\n")
        f.write(f"RR_bpm    : {_fmt_result(state.rr)}\n")


# ── Completion + wait loop ─────────────────────────────────────────────────────

def _session_complete() -> bool:
    return bool(state.error) or state.inf_done


async def _wait_for_completion(timeout: float) -> bool:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        await asyncio.sleep(0.05)
        if _session_complete():
            return True
    return False


def _save_session(rec_num: int) -> None:
    pstats.end_session()
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    base = _base_name(rec_num, ts)

    print()
    print(f"  ── Result  (Subject: {subject_code}  |  Activity: {activity}  |  Rec #{rec_num:03d}) ──")
    print(f"  HR : {_fmt_result(state.hr)} BPM")
    print(f"  RR : {_fmt_result(state.rr)} breaths/min")

    # WAV
    if state.audio_samples and state.expected_bytes > 0:
        wav_path = output_dir / f"{base}.wav"
        save_wav(bytearray(state.audio_samples), wav_path)
        kb  = wav_path.stat().st_size / 1024
        dur = len(state.audio_samples) / 2 / SAMPLE_RATE
        print(f"  WAV: {wav_path}  ({kb:.1f} kB, {dur:.2f} s)")
    else:
        print("  [WARN  ] No audio received — WAV not saved.")

    # TXT sidecar
    txt_path = output_dir / f"{base}.txt"
    save_results_txt(txt_path)
    print(f"  TXT: {txt_path}")

    # Excel log
    notes = ""
    if pstats.audio_pkts_lost > 0:
        notes = f"{pstats.audio_pkts_lost} pkt(s) lost ({pstats.audio_loss_pct:.2f}% loss)"
    if state.error:
        notes = f"Firmware error: {state.error}"

    _append_excel_row(
        rec_num, state.hr, state.rr,
        pstats.audio_pkts_lost, pstats.audio_loss_pct,
        notes,
    )
    print(f"  LOG: {_excel_log_path()}")
    print()


# ── Main ───────────────────────────────────────────────────────────────────────

async def keepalive_loop(client):
    while True:
        try:
            if idle_keepalive:
                await client.write_gatt_char(NUS_RX_CHAR_UUID, b"KA",
                                             response=False)
        except Exception:
            pass
        await asyncio.sleep(0.25)


async def run():
    prompt_session_info()
    _init_excel_log()

    print(f"Scanning for '{DEVICE_NAME}' (timeout={SCAN_TIMEOUT:.0f} s)…")
    device = await BleakScanner.find_device_by_name(DEVICE_NAME, timeout=SCAN_TIMEOUT)
    if device is None:
        print(f"ERROR: '{DEVICE_NAME}' not found within {SCAN_TIMEOUT:.0f} s.")
        sys.exit(1)

    print(f"  Found: {DEVICE_NAME}  [{device.address}]")

    async with BleakClient(device) as client:
        mtu = getattr(client, "mtu_size", "?")
        print(f"Connected!  MTU = {mtu}")
        await client.start_notify(NUS_TX_CHAR_UUID, handle_notification)
        ka_task = asyncio.create_task(keepalive_loop(client))

        try:
            await asyncio.wait_for(ready_event.wait(), timeout=READY_TIMEOUT)
        except asyncio.TimeoutError:
            print(f"  [WARN  ] No READY within {READY_TIMEOUT:.0f} s — proceeding.")

        print()
        print("Commands:")
        print("  ENTER  — start recording (HR + RR)")
        print("  q      — quit")
        print()

        rec_num = 1
        try:
            while True:
                sep = "─" * 55
                inp = await asyncio.get_event_loop().run_in_executor(
                    None,
                    lambda: input(
                        f"{sep}\n"
                        f"[{subject_code} | {activity} | #{rec_num:03d}]  "
                        "ENTER=REC  q=quit: "
                    )
                )
                cmd = inp.strip().lower()

                if cmd == "q":
                    print("Quitting.")
                    break

                elif cmd == "t":
                    await client.write_gatt_char(NUS_RX_CHAR_UUID, b"TEST",
                                                 response=False)
                    continue

                # ── REC ──
                globals()["idle_keepalive"] = False
                state.reset()
                print("\nSending REC…")
                await client.write_gatt_char(NUS_RX_CHAR_UUID, b"REC",
                                             response=False)

                completed = await _wait_for_completion(REC_TIMEOUT)

                if not completed:
                    print("\n  [WARN  ] Timeout — saving whatever arrived.")
                elif state.error:
                    print(f"\n  [ERROR ] Firmware error: {state.error}")
                else:
                    print("\n  [DONE  ] REC complete.")

                _save_session(rec_num)
                rec_num += 1
                globals()["idle_keepalive"] = True

        except KeyboardInterrupt:
            print("\nInterrupted by user.")
        finally:
            ka_task.cancel()
            await client.stop_notify(NUS_TX_CHAR_UUID)
            print("Disconnected.")


if __name__ == "__main__":
    asyncio.run(run())
